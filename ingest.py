from __future__ import annotations

import base64
import re
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# Utilities for reading and preparing input files.
# - Locate input directories and list files
# - Convert files to data URLs for model upload
# - Extract text from Excel files and parse email bodies/attachments


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".bmp",
    ".gif",
    ".xlsx",
    ".xlsm",
    ".xls",
    ".eml",
}

SUPPORTED_ATTACHMENT_EXTENSIONS = {
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".bmp",
    ".gif",
    ".xlsx",
    ".xlsm",
    ".xls",
}


# Resolve the best input directory for processing.
def resolve_input_dir(raw_dir: str | None) -> Path:
    base_dir = Path(__file__).resolve().parent
    candidates: list[Path] = []
    if raw_dir:
        candidates.append(Path(raw_dir))
    candidates.extend([base_dir / "input", base_dir / "Input"])

    for candidate in candidates:
        if candidate.exists():
            return candidate.resolve()

    if raw_dir:
        return Path(raw_dir).resolve()
    return (base_dir / "input").resolve()


# List files in a directory recursively.
def list_input_files(input_dir: Path) -> list[Path]:
    if not input_dir.exists():
        return []
    files = [path for path in input_dir.rglob("*") if path.is_file()]
    return sorted(files, key=lambda path: str(path).lower())


# Convert a file into a browser-ready data URL.
def file_to_data_url(file_path: Path) -> str:
    extension = file_path.suffix.lower().lstrip(".")
    mime_map = {
        "pdf": "application/pdf",
        "png": "image/png",
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "tif": "image/tiff",
        "tiff": "image/tiff",
        "bmp": "image/bmp",
        "gif": "image/gif",
    }
    mime_type = mime_map.get(extension, "application/octet-stream")
    encoded = base64.b64encode(file_path.read_bytes()).decode("utf-8")
    return f"data:{mime_type};base64,{encoded}"


# Extract readable values from an Excel workbook.
def extract_excel_text(file_path: Path) -> str:
    if file_path.suffix.lower() == ".xls":
        raise ValueError("Legacy .xls files are not supported by openpyxl. Please convert to .xlsx first.")

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet_sections: list[str] = []

    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        rows_text: list[str] = [f"Sheet: {sheet.title}"]
        for row_index, row in enumerate(rows, start=1):
            cleaned_values = ["" if value is None else str(value).strip() for value in row]
            if not any(value for value in cleaned_values):
                continue
            rows_text.append(f"Row {row_index}: | " + " | ".join(cleaned_values))

        if len(rows_text) > 1:
            sheet_sections.append("\n".join(rows_text))

    workbook.close()
    return "\n\n".join(sheet_sections) if sheet_sections else "No rows found in spreadsheet."


# Pull out the visible text from an email body.
def parse_email_body(message) -> str:
    body_parts: list[str] = []
    for part in message.walk():
        if part.is_multipart():
            continue
        content_type = part.get_content_type()
        if content_type in {"text/plain", "text/html"}:
            payload = part.get_payload(decode=True)
            if not payload:
                continue
            text = payload.decode("utf-8", errors="replace")
            if text.strip():
                body_parts.append(text)
    return "\n\n".join(body_parts)


# Gather supported email attachments and metadata.
def parse_email_attachments(message) -> list[dict[str, Any]]:
    attachments: list[dict[str, Any]] = []
    for part in message.walk():
        filename = part.get_filename()
        if not filename:
            continue
        payload = part.get_payload(decode=True)
        if payload is None:
            continue
        suffix = Path(filename).suffix.lower()
        if suffix not in SUPPORTED_ATTACHMENT_EXTENSIONS:
            continue
        attachments.append({
            "filename": filename,
            "extension": suffix,
            "content_type": part.get_content_type(),
            "payload": payload,
        })
    return attachments


# Read a .eml file and return cleaned email fields.
def extract_email_metadata(file_path: Path) -> dict[str, Any]:
    raw = file_path.read_bytes()
    message = BytesParser(policy=policy.default).parsebytes(raw)
    body_text = parse_email_body(message)
    clean_body = re.sub(r"\s+", " ", body_text).strip()
    return {
        "subject": (message.get("Subject") or "").strip(),
        "from": (message.get("From") or "").strip(),
        "to": (message.get("To") or "").strip(),
        "body": clean_body,
        "attachments": parse_email_attachments(message),
    }


# Detect wording suggesting a missing attachment in the email.
def detect_missing_attachment(body: str) -> bool:
    if not body:
        return False
    hints = (
        r"attached",
        r"enclosed",
        r"please find attached",
        r"see attached",
        r"see enclosure",
        r"forwarding attached",
        r"attached herewith",
        r"as attached",
    )
    pattern = "|".join(hints)
    return bool(re.search(pattern, body, flags=re.IGNORECASE))

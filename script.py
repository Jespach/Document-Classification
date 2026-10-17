from __future__ import annotations

import argparse
import base64
import json
import os
from pathlib import Path
from typing import Any

import requests
from dotenv import load_dotenv
from openpyxl import load_workbook


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".bmp",
    ".gif",
    ".xlsx",
    ".xlsm",
    ".xls",
}


def resolve_input_dir(raw_dir: str | None) -> Path:
    base_dir = Path(__file__).resolve().parent
    candidates = []
    if raw_dir:
        candidates.append(Path(raw_dir))
    candidates.extend([
        base_dir / "input",
        base_dir / "Input",
    ])

    for candidate in candidates:
        if candidate.exists():
            return candidate.resolve()

    if raw_dir:
        return Path(raw_dir).resolve()
    return (base_dir / "input").resolve()


def list_input_files(input_dir: Path) -> list[Path]:
    if not input_dir.exists():
        return []
    files = [path for path in input_dir.rglob("*") if path.is_file()]
    return sorted(files, key=lambda path: str(path).lower())


def get_api_key() -> str | None:
    load_dotenv()
    return os.getenv("OPENROUTER_API_KEY") or os.getenv("API_KEY")


JSON_SCHEMA = {
    "type": "object",
    "properties": {
        "label": {
            "type": "string",
            "enum": [
                "supplier_invoice",
                "credit_note",
                "statement",
                "remittance_advice",
                "other",
            ],
        },
        "confidence": {
            "type": "number",
            "minimum": 0,
            "maximum": 1,
        },
        "evidence": {
            "type": "array",
            "items": {"type": "string"},
            "minItems": 1,
        },
        "reasoning": {"type": "string"},
    },
    "required": ["label", "confidence", "evidence", "reasoning"],
    "additionalProperties": False,
}


def sanitize_classification_payload(payload: dict[str, Any]) -> dict[str, Any]:
    label = payload.get("label")
    if label not in {"supplier_invoice", "credit_note", "statement", "remittance_advice", "other"}:
        raise ValueError(f"Invalid label returned by model: {label!r}")

    confidence = payload.get("confidence")
    try:
        confidence_value = float(confidence)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid confidence returned by model: {confidence!r}") from exc

    if confidence_value < 0 or confidence_value > 1:
        raise ValueError(f"Confidence must be between 0 and 1, got {confidence_value}")

    evidence = payload.get("evidence")
    if not isinstance(evidence, list) or not evidence or not all(isinstance(item, str) for item in evidence):
        raise ValueError("Evidence must be a non-empty list of strings.")

    reasoning = payload.get("reasoning")
    if not isinstance(reasoning, str) or not reasoning.strip():
        raise ValueError("Reasoning must be a non-empty string.")

    return {
        "label": label,
        "confidence": round(confidence_value, 4),
        "evidence": [str(item).strip() for item in evidence if str(item).strip()],
        "reasoning": reasoning.strip(),
    }


def load_prompt_template() -> str:
    prompt_path = Path(__file__).resolve().parent / "prompt.md"
    if prompt_path.exists():
        template = prompt_path.read_text(encoding="utf-8")
        return template.replace("{{DOCUMENT_TEXT}}", "").strip()
    return "" 


def build_document_prompt(file_path: Path) -> str:
    prompt_template = load_prompt_template()
    if prompt_template:
        return prompt_template

    return (
        "You are a document classification system used by a finance team to categorise incoming supplier-related documents. "
        "Review the supplied document and return only a valid JSON object matching this schema: "
        "{\"label\": \"supplier_invoice|credit_note|statement|remittance_advice|other\", "
        "\"confidence\": 0.0, \"evidence\": [\"specific observation\"], \"reasoning\": \"brief explanation\"}. "
        "Do not include markdown or extra text."
    )


def file_to_data_url(file_path: Path) -> str:
    extension = file_path.suffix.lower().lstrip(".")

    mime_map = {
        "pdf": "application/pdf",
        "png": "image/png",
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "tif": "image/tiff",
        "tiff": "image/tiff",
        "bmp": "image/bmp",
        "gif": "image/gif",
    }

    mime_type = mime_map.get(extension, "application/octet-stream")
    encoded = base64.b64encode(file_path.read_bytes()).decode("utf-8")
    return f"data:{mime_type};base64,{encoded}"


def extract_excel_text(file_path: Path) -> str:
    if file_path.suffix.lower() == ".xls":
        raise ValueError("Legacy .xls files are not supported by openpyxl. Please convert to .xlsx first.")

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet_sections: list[str] = []

    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        rows_text: list[str] = [f"Sheet: {sheet.title}"]
        for row_index, row in enumerate(rows, start=1):
            cleaned_values = ["" if value is None else str(value).strip() for value in row]
            if not any(value for value in cleaned_values):
                continue
            rows_text.append(f"Row {row_index}: | " + " | ".join(cleaned_values))

        if len(rows_text) > 1:
            sheet_sections.append("\n".join(rows_text))

    workbook.close()
    return "\n\n".join(sheet_sections) if sheet_sections else "No rows found in spreadsheet."


def build_user_message(file_path: Path) -> dict[str, Any]:
    prompt = build_document_prompt(file_path)
    ext = file_path.suffix.lower()

    if ext in {".xlsx", ".xlsm", ".xls"}:
        spreadsheet_text = extract_excel_text(file_path)
        return {
            "role": "user",
            "content": [
                {
                    "type": "text",
                    "text": f"{prompt}\n\nStructured spreadsheet content:\n{spreadsheet_text}",
                }
            ],
        }

    data_url = file_to_data_url(file_path)
    if ext == ".pdf":
        return {
            "role": "user",
            "content": [
                {"type": "text", "text": prompt},
                {
                    "type": "file",
                    "file": {
                        "filename": file_path.name,
                        "file_data": data_url,
                    },
                },
            ],
        }

    return {
        "role": "user",
        "content": [
            {"type": "text", "text": prompt},
            {
                "type": "image_url",
                "image_url": {"url": data_url},
            },
        ],
    }


def get_openrouter_base_url() -> str:
    load_dotenv()
    return os.getenv("OPENROUTER_BASE_URL") or "https://openrouter.ai/api/v1"


def call_openrouter(messages: list[dict[str, Any]], api_key: str) -> dict[str, Any]:
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }

    payload = {
        "model": "openai/gpt-5.6-luna",
        "messages": messages,
        "reasoning": {"enabled": True},
        "response_format": {
            "type": "json_schema",
            "json_schema": {
                "name": "document_classification",
                "schema": JSON_SCHEMA,
                "strict": True,
            },
        },
    }

    response = requests.post(
        url=f"{get_openrouter_base_url()}/chat/completions",
        headers=headers,
        json=payload,
        timeout=120,
    )
    response.raise_for_status()
    return response.json()


def extract_cost(response_json: dict[str, Any]) -> float | None:
    usage = response_json.get("usage") or {}
    if not isinstance(usage, dict):
        return None

    cost_candidates = [
        usage.get("cost"),
        usage.get("cost_usd"),
        usage.get("total_cost"),
        usage.get("total_cost_usd"),
    ]

    for candidate in cost_candidates:
        if candidate is None:
            continue
        try:
            return float(candidate)
        except (TypeError, ValueError):
            continue

    return None


def format_table_row(file_name: str, document_type: str, reasoning: str, cost: float | None) -> str:
    clean_reasoning = reasoning.strip().replace("\n", " ")
    if len(clean_reasoning) > 80:
        clean_reasoning = clean_reasoning[:77] + "..."

    cost_label = f"${cost:.6f}" if cost is not None else "n/a"
    return f"{file_name:<30} | {document_type:<18} | {clean_reasoning:<80} | {cost_label}"


def parse_structured_response(raw_content: str | None) -> dict[str, Any]:
    if not raw_content:
        raise ValueError("OpenRouter returned an empty response.")

    cleaned = raw_content.strip()
    if cleaned.startswith("```"):
        cleaned = cleaned.strip("`")
        if cleaned.lower().startswith("json"):
            cleaned = cleaned[4:].lstrip()

    try:
        parsed = json.loads(cleaned)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Model response was not valid JSON: {cleaned[:200]}") from exc

    if not isinstance(parsed, dict):
        raise ValueError("Model response was not a JSON object.")

    return sanitize_classification_payload(parsed)


def analyze_document(file_path: Path, api_key: str) -> dict[str, Any]:
    first_message = build_user_message(file_path)
    first_response = call_openrouter([first_message], api_key)
    first_cost = extract_cost(first_response)
    assistant_message = first_response["choices"][0]["message"]

    follow_up_messages = [
        first_message,
        {
            "role": "assistant",
            "content": assistant_message.get("content"),
            "reasoning_details": assistant_message.get("reasoning_details"),
        },
        {
            "role": "user",
            "content": "Return the final answer strictly as valid JSON matching the required schema only.",
        },
    ]

    final_response = call_openrouter(follow_up_messages, api_key)
    final_cost = extract_cost(final_response)
    final_message = final_response["choices"][0]["message"]
    structured = parse_structured_response(final_message.get("content"))

    total_cost = None
    if first_cost is not None and final_cost is not None:
        total_cost = first_cost + final_cost
    elif first_cost is not None:
        total_cost = first_cost
    elif final_cost is not None:
        total_cost = final_cost

    return {
        "file": file_path.name,
        "result": structured,
        "cost_usd": total_cost,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Ingest supported documents and send each file through OpenRouter for analysis.")
    parser.add_argument(
        "--input-dir",
        default=None,
        help="Directory containing documents to analyze. Defaults to ./input or ./Input if present.",
    )
    args = parser.parse_args()

    input_dir = resolve_input_dir(args.input_dir)
    if not input_dir.exists():
        print(f"Input directory does not exist: {input_dir}")
        return

    supported_files = []
    for file_path in list_input_files(input_dir):
        if file_path.suffix.lower() in SUPPORTED_EXTENSIONS:
            supported_files.append(file_path)

    if not supported_files:
        print(f"No supported image or PDF files found in {input_dir}")
        return

    api_key = get_api_key()
    if not api_key:
        print("No OpenRouter API key found. Add OPENROUTER_API_KEY or API_KEY to your .env file.")
        return

    print(f"Found {len(supported_files)} supported file(s). Processing...")
    print(f"{'File':<30} | {'Document Type':<18} | {'Reasoning':<80} | {'Cost (USD)'}")
    print("-" * 150)
    for file_path in supported_files:
        try:
            result = analyze_document(file_path, api_key)
            row = format_table_row(
                result["file"],
                result["result"].get("label", "unknown"),
                result["result"].get("reasoning", "No reasoning returned."),
                result.get("cost_usd"),
            )
            print(row)
        except Exception as exc:
            print(format_table_row(file_path.name, "error", str(exc), None))


if __name__ == "__main__":
    main()
